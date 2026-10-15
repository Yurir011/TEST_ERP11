"""견적서/거래명세서를 Doc 폴더에서 옮겨온 실제 양식(backend/app/assets/templates)의 빈 칸에
채워 넣어 엑셀을 만든다. 임의로 새 워크북을 그리지 않고, 서식이 이미 갖춰진 원본 파일을 그대로 사용한다."""
import os
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from sqlalchemy.orm import Session

from app.config import settings
from app.logging_config import get_logger
from app.models.project_document import DOC_CURRENCY_SYMBOLS, DocCurrency, ProjectDocType, ProjectDocument, calc_amount

logger = get_logger("ProjectDocumentTemplate")

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "assets" / "templates"

DOC_NO_PLACEHOLDER = "XXXX-XX"

# 단가/합계/소계/부가세/합계 칸은 통화 기호를 숫자 바로 앞에 붙여 표시한다.
TOTAL_MONEY_FORMATS: dict[str, str] = {c.value: f'"{DOC_CURRENCY_SYMBOLS[c]}"#,##0' for c in DocCurrency}

# 상단 "일금 ___ 원정" 줄의 "원정"은 통화 기호가 아니라 고정 텍스트라 서식만으론 안 바뀐다. 통화별 단어로 직접 바꿔준다.
CURRENCY_SUFFIX_WORD: dict[str, str] = {
    DocCurrency.KRW.value: "원정",
    DocCurrency.USD.value: "불정",
    DocCurrency.JPY.value: "엔정",
}

# 소계/부가세/합계 줄은 글자가 없는 서체(HY그래픽M)를 쓰고 있어 ¥ 같은 기호가 네모로 깨진다.
# 통화 서식을 적용할 때는 글꼴도 화폐 기호를 지원하는 서체로 함께 바꿔준다.
MONEY_FONT_NAME = "맑은 고딕"


def _apply_money_format(cell, number_format: str, size: float | None = None) -> None:
    cell.number_format = number_format
    f = cell.font
    cell.font = Font(
        name=MONEY_FONT_NAME, size=size if size is not None else f.sz, bold=f.bold, italic=f.italic, color=f.color, underline=f.underline
    )


def _set_horizontal(cell, horizontal: str) -> None:
    a = cell.alignment
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=a.vertical,
        wrap_text=a.wrap_text,
        shrink_to_fit=a.shrink_to_fit,
        indent=a.indent,
    )


def _fit_to_page_width(ws) -> None:
    """양식을 A4 세로 한 장의 좌우 폭에 맞춰 출력하도록 지정한다. 폭을 기준으로 배율을 줄이고 남는 부분은 여백으로 둔다.
    프린터/Excel 설정이 달라도 오른쪽이 잘리거나 다음 장으로 넘어가지 않도록, 내용(글자/테두리/이미지)이 있는 범위만 인쇄 영역으로 잡는다."""
    last_row = last_col = 1
    for row in ws.iter_rows():
        for c in row:
            b = c.border
            if c.value is not None or b.left.style or b.right.style or b.top.style or b.bottom.style:
                last_row, last_col = max(last_row, c.row), max(last_col, c.column)
    for rng in ws.merged_cells.ranges:
        last_row, last_col = max(last_row, rng.max_row), max(last_col, rng.max_col)
    # 직인/로고 이미지가 내용 영역 바깥 열까지 걸쳐 있으면 그 열까지 인쇄 영역에 포함한다. (양식 이미지는 시작/끝 셀이 지정된 방식이다)
    for img in ws._images:
        end = getattr(img.anchor, "to", None)
        if end is not None:
            last_col = max(last_col, end.col + 1)
    ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.scale = None
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # 높이는 자동 (폭 기준으로만 맞춘다)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    logger.debug(f"[ProjectDocumentTemplate] 인쇄 영역 지정: A1:{get_column_letter(last_col)}{last_row}, 폭 1장에 맞춤")


@dataclass(frozen=True)
class DocLayout:
    template_filename: str
    date_cell: str
    client_cell: str
    manager_cell: str
    doc_no_cell: str
    top_total_cell: str  # 상단 "합계 / 일금 ___원정" 줄의 금액 칸 (총액을 한 번 더 보여줌)
    total_words_cell: str  # 상단 줄의 "원정" 글자 칸
    item_start_row: int
    item_end_row: int
    subtotal_row: int
    vat_row: int
    total_row: int


# 두 양식은 회사 정보(사업자번호/주소/대표자 등)가 이미 고정 텍스트로 박혀있어 채울 필요가 없고,
# 행 구성만 서로 다르다 (열 의미는 공통: A=번호, B=내용, E=수량, F=단가, G=합계, I=비고).
LAYOUTS: dict[ProjectDocType, DocLayout] = {
    ProjectDocType.quotation: DocLayout(
        template_filename="견적서.xlsx",
        date_cell="H3",
        client_cell="H4",
        manager_cell="H5",
        doc_no_cell="H6",
        top_total_cell="G14",
        total_words_cell="F14",
        item_start_row=17,
        item_end_row=28,
        subtotal_row=29,
        vat_row=30,
        total_row=31,
    ),
    ProjectDocType.statement: DocLayout(
        template_filename="거래명세서.xlsx",
        date_cell="H2",
        client_cell="H3",
        manager_cell="H4",
        doc_no_cell="H5",
        top_total_cell="G12",
        total_words_cell="F12",
        item_start_row=14,
        item_end_row=25,
        subtotal_row=26,
        vat_row=27,
        total_row=28,
    ),
}


def _append_value(current: object, value: str) -> str:
    """'DATE : ' 같은 라벨 뒤에 값을 이어 붙인다. (템플릿 셀에 라벨+값이 한 셀에 들어있음)"""
    label = str(current or "").rstrip()
    return f"{label} {value}" if value else label


TAX_INVOICE_TEMPLATE = "세금계산서.xlsx"
TAX_INVOICE_ITEM_ROWS = (11, 12, 13, 14)  # 양식의 품목 칸은 4줄
# 공급자(왼쪽)/공급받는자(오른쪽) 칸 위치. 양식에는 공급자 칸에 우리 회사 정보가 이미 적혀 있으므로, 받은 세금계산서(매입)는
# 공급자 칸에 거래처, 공급받는자 칸에 우리 회사를 적도록 양쪽 모두 덮어쓴다.
_SUPPLIER_CELLS = {"reg_no": "E2", "name": "E3", "ceo": "J3", "address": "E4", "biz_type": "E5", "biz_class": "J5", "email": "E6"}
_RECIPIENT_CELLS = {"reg_no": "O2", "name": "O3", "ceo": "R3", "address": "O4", "biz_type": "O5", "biz_class": "R5", "email": "O6"}


def _party_values(name: str | None, reg_no: str | None, ceo: str | None, address: str | None, biz_type: str | None,
                  biz_class: str | None, email: str | None) -> dict[str, str]:
    return {
        "reg_no": reg_no or "", "name": name or "", "ceo": ceo or "", "address": address or "",
        "biz_type": biz_type or "", "biz_class": biz_class or "", "email": email or "",
    }


def _fill_tax_invoice_template(doc: ProjectDocument) -> str:
    """세금계산서: 홈택스 전자세금계산서와 같은 양식(A5 가로 한 장)을 채운다."""
    template_path = TEMPLATE_DIR / TAX_INVOICE_TEMPLATE
    logger.debug(f"[ProjectDocumentTemplate] 세금계산서 양식 채우기 시작: doc_id={doc.id}, direction={doc.direction}")
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # 우리 회사 정보는 양식의 공급자 칸에 원본 그대로 적혀 있는 값을 쓴다 (양쪽 칸을 덮어쓰기 전에 먼저 읽어둔다).
    ours = {key: str(ws[cell].value or "").strip() for key, cell in _SUPPLIER_CELLS.items()}
    # 상대방 정보: 거래처관리에 등록된 거래처의 값을 우선 쓰고, 비어 있는 칸은 사진에서 인식해 문서에 저장해 둔 값으로 채운다.
    client = doc.effective_client
    snap = doc.counterparty_info or {}

    def pick(client_value: str | None, key: str) -> str | None:
        return client_value or snap.get(key)

    theirs = _party_values(
        (client.name if client else None) or snap.get("name") or doc.client_name,
        pick(client.biz_reg_no if client else None, "reg_no"),
        pick(client.ceo_name if client else None, "ceo_name"),
        pick(client.address if client else None, "address"),
        pick(client.biz_type if client else None, "biz_type"),
        pick(client.biz_class if client else None, "biz_class"),
        pick(client.email if client else None, "email"),
    )
    is_sales = doc.direction == "sales" if doc.direction else doc.purpose_type == "청구"
    supplier, recipient = (ours, theirs) if is_sales else (theirs, ours)
    for key, cell in _SUPPLIER_CELLS.items():
        ws[cell] = supplier[key] or None
    for key, cell in _RECIPIENT_CELLS.items():
        ws[cell] = recipient[key] or None

    # 공급받는자 칸의 글꼴/정렬/이메일 칸 구성을 공급자 칸 기준으로 맞춘다 (양식 원본은 양쪽이 서로 다르게 되어 있음).
    for key, s_cell in _SUPPLIER_CELLS.items():
        s, r = ws[s_cell], ws[_RECIPIENT_CELLS[key]]
        r.font = Font(name=s.font.name, size=s.font.sz, bold=s.font.b, italic=s.font.i, color=r.font.color)
        r.alignment = Alignment(horizontal=s.alignment.horizontal, vertical=s.alignment.vertical, wrap_text=s.alignment.wrap_text)
    # 공급자는 이메일이 두 줄을 합친 한 칸이므로 공급받는자도 두 줄(6~7행)을 하나로 합친다.
    ws["O7"] = None
    ws.merge_cells("N6:N7")
    ws.unmerge_cells("O6:S6")
    ws.unmerge_cells("O7:S7")
    ws.merge_cells("O6:S7")
    logger.debug(f"[ProjectDocumentTemplate] 공급받는자 칸 서식을 공급자 기준으로 통일: doc_id={doc.id}")

    ws["O1"] = doc.approval_no

    # 품목: 양식 칸(4줄)보다 많으면 마지막 줄에 나머지를 "외 N건"으로 합산해 금액이 빠지지 않게 한다.
    rows_capacity = len(TAX_INVOICE_ITEM_ROWS)
    items = list(doc.items)
    lines: list[tuple[str, float | None, float | None, int, int, str | None]] = []
    for it in items[: rows_capacity - 1] if len(items) > rows_capacity else items:
        amount = calc_amount(it.quantity, it.unit_price)
        lines.append((it.content, it.quantity, it.unit_price, amount, round(amount * 0.1), it.note))
    if len(items) > rows_capacity:
        rest = items[rows_capacity - 1:]
        amount = sum(calc_amount(it.quantity, it.unit_price) for it in rest)
        lines.append((f"외 {len(rest)}건", None, None, amount, round(amount * 0.1), None))

    subtotal = sum(calc_amount(it.quantity, it.unit_price) for it in items)
    vat = round(subtotal * 0.1)
    total = subtotal + vat

    ws["A9"] = doc.issue_date.isoformat()
    ws["E9"] = subtotal
    ws["H9"] = vat
    ws["M9"] = "해당없음"
    ws["E9"].number_format = ws["H9"].number_format = "#,##0"

    for row, line in zip(TAX_INVOICE_ITEM_ROWS, lines + [None] * (rows_capacity - len(lines))):
        if line is None:
            continue
        content, qty, unit_price, amount, line_vat, note = line
        ws[f"A{row}"] = f"{doc.issue_date.month:02d}"
        ws[f"C{row}"] = f"{doc.issue_date.day:02d}"
        ws[f"D{row}"] = content
        ws[f"K{row}"] = qty
        ws[f"M{row}"] = unit_price
        ws[f"O{row}"] = amount
        ws[f"P{row}"] = line_vat
        ws[f"I{row}"] = note  # 품목의 비고(사진에서 읽은 규격 포함)는 양식의 규격 칸에 적는다 (홈택스 원본도 비고 칸은 비워둔다)
        if unit_price is not None:
            ws[f"M{row}"].number_format = "#,##0" if float(unit_price).is_integer() else "#,##0.0#"
        ws[f"K{row}"].number_format = "#,##0"
        ws[f"O{row}"].number_format = ws[f"P{row}"].number_format = "#,##0"

    # 합계금액 줄: 대금을 이미 받은/지급한 건은 현금, 아직이면 외상미수금 칸에 합계를 적는다.
    ws["A16"] = total
    ws["A16"].number_format = "#,##0"
    paid = doc.payment_received_hint is True
    ws["F16" if paid else "O16"] = total
    ws["F16" if paid else "O16"].number_format = "#,##0"
    ws["P15"] = f"이 금액을 ( {doc.purpose_type} ) 함"

    # 양식 맨 아래의 "공급자 정보만 기재되어 있습니다" 안내 문구는 채워진 문서에는 필요 없어 지우고, 출력 범위에서도 뺀다.
    ws["A18"] = None
    ws.print_area = "A1:S16"
    # A5 가로 한 장에 딱 맞게 출력되도록 용지/배율/여백을 지정한다.
    ws.page_setup.paperSize = ws.PAPERSIZE_A5
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.2
    ws.page_margins.header = ws.page_margins.footer = 0
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = True

    os.makedirs(settings.project_documents_dir, exist_ok=True)
    out_path = os.path.join(settings.project_documents_dir, f"{doc.id}_{doc.doc_type.value}.xlsx")
    wb.save(out_path)
    logger.debug(f"[ProjectDocumentTemplate] 세금계산서 양식 채우기 완료: doc_id={doc.id}, file={out_path}")
    return out_path


PURCHASE_ORDER_TEMPLATE = "발주서.xlsx"
# 발주서 양식의 칸 위치 (A=번호, B~C=내용(병합), D=수량, E=단가, F~G=합계(병합), H=비고). 회사 정보/직인/로고는 양식에 박혀 있어 채우지 않는다.
PO_DATE_CELL, PO_CLIENT_CELL, PO_MANAGER_CELL = "G2", "G3", "G4"
PO_TOP_TOTAL_CELL, PO_TOTAL_WORDS_CELL = "F12", "E12"
PO_TOP_WORDS_CELL = "C12"  # "일금 ___" 한글 금액 칸
PO_ITEM_START_ROW, PO_ITEM_END_ROW = 14, 27
PO_SUBTOTAL_CELL, PO_VAT_CELL, PO_TOTAL_CELL = "F28", "F29", "F30"
PO_REMARK_CELL = "B31"
PO_END_MARK = " 이 하 여 백"


def _fill_purchase_order_template(doc: ProjectDocument) -> str:
    """발주서: 양식(Doc/발주서.xlsx)의 건별 입력 칸(날짜/거래처/담당자/품목/REMARK)만 채우고 나머지는 양식 그대로 둔다."""
    template_path = TEMPLATE_DIR / PURCHASE_ORDER_TEMPLATE
    logger.debug(f"[ProjectDocumentTemplate] 발주서 양식 채우기 시작: doc_id={doc.id}, template={template_path}")
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    d = doc.issue_date
    ws[PO_DATE_CELL] = f"DATE : {d.year}년 {d.month:02d}월 {d.day:02d}일"
    ws[PO_CLIENT_CELL] = f"거래처명 : {doc.client_name}"
    ws[PO_MANAGER_CELL] = f"담당자 : {doc.manager_name}" if doc.manager_name else "담당자 :"

    money_format = TOTAL_MONEY_FORMATS[doc.currency]
    items = list(doc.items)
    capacity = PO_ITEM_END_ROW - PO_ITEM_START_ROW + 1
    if len(items) > capacity:
        logger.debug(f"[ProjectDocumentTemplate] 발주서 양식 항목 칸({PO_ITEM_END_ROW}행)을 초과해 나머지는 생략: doc_id={doc.id}")
        items = items[:capacity]

    # 양식에 예시로 적혀 있는 품목/"이 하 여 백"을 모두 지우고 시작한다 (병합 칸의 오른쪽 셀은 건드리지 않는다).
    for row in range(PO_ITEM_START_ROW, PO_ITEM_END_ROW + 1):
        for col in "ABDEFH":
            ws[f"{col}{row}"] = None

    subtotal = 0
    for offset, item in enumerate(items):
        row = PO_ITEM_START_ROW + offset
        amount = calc_amount(item.quantity, item.unit_price)
        subtotal += amount
        ws[f"A{row}"] = offset + 1
        ws[f"B{row}"] = item.content
        ws[f"D{row}"] = item.quantity
        ws[f"E{row}"] = item.unit_price
        ws[f"F{row}"] = amount
        if item.note:
            ws[f"H{row}"] = item.note
    for row in range(PO_ITEM_START_ROW, PO_ITEM_END_ROW + 1):
        _apply_money_format(ws[f"E{row}"], money_format)
        _apply_money_format(ws[f"F{row}"], money_format)
        _set_horizontal(ws[f"B{row}"], "center")
        _set_horizontal(ws[f"H{row}"], "center")
        _set_horizontal(ws[f"F{row}"], "right")
        # 단가 칸은 양식 폭이 좁아(약 7자) 큰 금액은 글자를 줄여서 넣는다.
        e = ws[f"E{row}"]
        e.alignment = Alignment(horizontal="right", vertical=e.alignment.vertical, shrink_to_fit=True)
    end_row = PO_ITEM_START_ROW + len(items)
    if end_row <= PO_ITEM_END_ROW:
        ws[f"B{end_row}"] = PO_END_MARK

    vat = round(subtotal * 0.1)
    ws[PO_SUBTOTAL_CELL] = subtotal
    ws[PO_VAT_CELL] = vat
    ws[PO_TOTAL_CELL] = subtotal + vat
    for ref in (PO_SUBTOTAL_CELL, PO_VAT_CELL, PO_TOTAL_CELL, PO_TOP_TOTAL_CELL):
        _apply_money_format(ws[ref], money_format, size=ws[PO_SUBTOTAL_CELL].font.sz if ref == PO_TOTAL_CELL else None)
        _set_horizontal(ws[ref], "right")
    ws[PO_TOTAL_WORDS_CELL] = CURRENCY_SUFFIX_WORD[doc.currency]
    # "일금 ___ 원정" 칸은 한글 숫자 서식이라 금액이 크면 '###'으로 깨진다. 칸에 맞춰 글자를 줄여 항상 보이게 한다.
    c = ws[PO_TOP_WORDS_CELL]
    c.alignment = Alignment(horizontal=c.alignment.horizontal, vertical=c.alignment.vertical, shrink_to_fit=True)

    ws[PO_REMARK_CELL] = doc.remark or None
    _fit_to_page_width(ws)

    os.makedirs(settings.project_documents_dir, exist_ok=True)
    out_path = os.path.join(settings.project_documents_dir, f"{doc.id}_{doc.doc_type.value}.xlsx")
    wb.save(out_path)
    logger.debug(f"[ProjectDocumentTemplate] 발주서 양식 채우기 완료: doc_id={doc.id}, items={len(items)}, file={out_path}")
    return out_path


def fill_project_document_template(doc: ProjectDocument) -> str:
    """양식 원본을 열어 빈 칸을 채우고 storage에 저장한 뒤, 저장 경로를 반환한다."""
    if doc.doc_type == ProjectDocType.tax_invoice:
        return _fill_tax_invoice_template(doc)
    if doc.doc_type == ProjectDocType.purchase_order:
        return _fill_purchase_order_template(doc)
    layout = LAYOUTS[doc.doc_type]

    template_path = TEMPLATE_DIR / layout.template_filename
    logger.debug(f"[ProjectDocumentTemplate] 양식 채우기 시작: doc_id={doc.id}, template={template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    ws[layout.date_cell] = _append_value(ws[layout.date_cell].value, doc.issue_date.isoformat())
    ws[layout.client_cell] = _append_value(ws[layout.client_cell].value, doc.client_name)
    ws[layout.manager_cell] = _append_value(ws[layout.manager_cell].value, doc.manager_name or "")

    if doc.doc_no:
        current = str(ws[layout.doc_no_cell].value or "")
        ws[layout.doc_no_cell] = current.replace(DOC_NO_PLACEHOLDER, doc.doc_no)
    else:
        # 거래명세서는 견적번호가 필요 없어, 양식에 남아있는 안내 문구(플레이스홀더)를 비워둔다.
        ws[layout.doc_no_cell] = None

    total_money_format = TOTAL_MONEY_FORMATS[doc.currency]

    subtotal = 0
    item_count = min(len(doc.items), layout.item_end_row - layout.item_start_row + 1)
    if len(doc.items) > item_count:
        logger.debug(f"[ProjectDocumentTemplate] 양식 항목 칸({layout.item_end_row}행)을 초과해 나머지는 생략: doc_id={doc.id}")

    for offset in range(layout.item_end_row - layout.item_start_row + 1):
        row = layout.item_start_row + offset
        if offset < item_count:
            item = doc.items[offset]
            amount = calc_amount(item.quantity, item.unit_price)
            subtotal += amount
            ws[f"B{row}"] = item.content
            ws[f"E{row}"] = item.quantity
            ws[f"F{row}"] = item.unit_price
            ws[f"G{row}"] = amount
            if item.note:
                ws[f"I{row}"] = item.note
        else:
            # 사용하지 않는 행은 남아있는 수식(예: =E*F)을 지워 0원으로 표시되지 않게 한다.
            ws[f"G{row}"] = None
        # 양식 원본의 회계 서식은 기호가 칸 왼쪽 끝에 떨어져 보여서, 원화 포함 모두 기호를 숫자 바로 앞에 붙이는 서식으로 덮어쓴다.
        _apply_money_format(ws[f"F{row}"], total_money_format)
        _apply_money_format(ws[f"G{row}"], total_money_format)
        # 내용/비고는 가운데 맞춤, 단가/합계(금액)는 오른쪽 맞춤 - 양식 원본은 행마다 정렬이 들쭉날쭉해 통일한다.
        _set_horizontal(ws[f"B{row}"], "center")
        _set_horizontal(ws[f"I{row}"], "center")
        _set_horizontal(ws[f"F{row}"], "right")
        _set_horizontal(ws[f"G{row}"], "right")

    vat = round(subtotal * 0.1)
    total = subtotal + vat
    ws[f"G{layout.subtotal_row}"] = subtotal
    ws[f"G{layout.vat_row}"] = vat
    ws[f"G{layout.total_row}"] = total
    for cell_ref in (f"G{layout.subtotal_row}", f"G{layout.vat_row}", f"G{layout.total_row}", layout.top_total_cell):
        cell = ws[cell_ref]
        # 양식 원본의 최종 합계(세금 포함) 칸만 9pt라 공급가액/부가세(10pt)보다 작아 보여, 같은 크기로 맞춘다.
        _apply_money_format(
            cell, total_money_format, size=ws[f"G{layout.subtotal_row}"].font.sz if cell_ref == f"G{layout.total_row}" else None
        )
        _set_horizontal(cell, "right")
    ws[layout.total_words_cell] = CURRENCY_SUFFIX_WORD[doc.currency]
    # "일금 ___ 원정" 칸(C열)은 한글 숫자 서식이라 금액이 크면 '###'으로 깨진다. 칸에 맞춰 글자를 줄여 항상 보이게 한다.
    words = ws[f"C{ws[layout.top_total_cell].row}"]
    words.alignment = Alignment(horizontal=words.alignment.horizontal, vertical=words.alignment.vertical, shrink_to_fit=True)
    _fit_to_page_width(ws)

    os.makedirs(settings.project_documents_dir, exist_ok=True)
    out_path = os.path.join(settings.project_documents_dir, f"{doc.id}_{doc.doc_type.value}.xlsx")
    wb.save(out_path)
    logger.debug(f"[ProjectDocumentTemplate] 양식 채우기 완료: doc_id={doc.id}, file={out_path}")
    return out_path


def next_doc_no(db: Session, doc_type: ProjectDocType, issue_date: date) -> str:
    """같은 날, 같은 문서 종류 기준으로 발행 순번을 매겨 MMDD-NN 형태로 반환한다."""
    count = (
        db.query(ProjectDocument)
        .filter(ProjectDocument.doc_type == doc_type, ProjectDocument.issue_date == issue_date)
        .count()
    )
    seq = count + 1
    doc_no = f"{issue_date.strftime('%m%d')}-{seq:02d}"
    logger.debug(f"[ProjectDocumentTemplate] 견적번호 채번: doc_type={doc_type}, issue_date={issue_date}, doc_no={doc_no}")
    return doc_no
