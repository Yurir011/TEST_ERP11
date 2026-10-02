from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.client import Client

HEADERS = [
    "업체명",
    "사업자번호",
    "대표자",
    "업태",
    "종목",
    "전화",
    "이메일",
    "주소",
    "은행",
    "계좌번호",
    "미수금",
    "미지급금",
    "담당자",
    "담당자 연락처",
    "담당자 이메일",
    "메모",
]
COLUMN_WIDTHS = [20, 16, 12, 16, 16, 16, 22, 28, 12, 20, 14, 14, 14, 16, 22, 24]


def _primary_contact(client: Client):
    return client.contacts[0] if client.contacts else None


def generate_client_list_excel(clients: list[Client]) -> bytes:
    """(검색 필터가 적용된) 거래처 목록을 엑셀 한 장으로 내보낸다. 담당자가 여러 명이면 대표 담당자 1명만 표시한다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "거래처목록"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    row_idx = 2
    for client in clients:
        contact = _primary_contact(client)
        contact_name = contact.name if contact else "-"
        if contact and len(client.contacts) > 1:
            contact_name += f" 외 {len(client.contacts) - 1}명"
        contact_phone = (contact.mobile_phone or contact.landline_phone) if contact else None

        ws.cell(row=row_idx, column=1, value=client.name)
        ws.cell(row=row_idx, column=2, value=client.biz_reg_no or "-")
        ws.cell(row=row_idx, column=3, value=client.ceo_name or "-")
        ws.cell(row=row_idx, column=4, value=client.biz_type or "-")
        ws.cell(row=row_idx, column=5, value=client.biz_class or "-")
        ws.cell(row=row_idx, column=6, value=client.phone or "-")
        ws.cell(row=row_idx, column=7, value=client.email or "-")
        ws.cell(row=row_idx, column=8, value=client.address or "-")
        ws.cell(row=row_idx, column=9, value=client.bank_name or "-")
        ws.cell(row=row_idx, column=10, value=client.bank_account or "-")
        for col, amount in ((11, client.receivable_amount), (12, client.payable_amount)):
            cell = ws.cell(row=row_idx, column=col, value=amount)
            cell.number_format = '"₩"#,##0'
        ws.cell(row=row_idx, column=13, value=contact_name)
        ws.cell(row=row_idx, column=14, value=contact_phone or "-")
        ws.cell(row=row_idx, column=15, value=contact.email if contact else "-")
        ws.cell(row=row_idx, column=16, value=client.memo or "-")
        row_idx += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
