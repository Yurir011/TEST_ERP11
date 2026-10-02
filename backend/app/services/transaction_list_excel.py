from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.transaction import Transaction, TransactionType

TYPE_LABELS = {TransactionType.sales: "매출", TransactionType.purchase: "매입"}

HEADERS = ["구분", "날짜", "거래처", "품목", "공급가액", "부가세", "합계", "세금계산서번호", "메모"]
COLUMN_WIDTHS = [8, 12, 20, 24, 14, 12, 14, 18, 24]


def generate_transaction_list_excel(transactions: list[Transaction]) -> bytes:
    """(검색/기간 필터가 적용된) 매입매출 내역 목록을 엑셀 한 장으로 내보낸다. 맨 아래에 매출/매입/차액 합계를 덧붙인다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매입매출내역"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    total_sales = 0
    total_purchase = 0
    row_idx = 2
    for tx in transactions:
        ws.cell(row=row_idx, column=1, value=TYPE_LABELS[tx.type])
        ws.cell(row=row_idx, column=2, value=tx.transaction_date.isoformat())
        ws.cell(row=row_idx, column=3, value=tx.client.name if tx.client else (tx.counterparty or "-"))
        ws.cell(row=row_idx, column=4, value=tx.item_name)
        for col, amount in ((5, tx.supply_amount), (6, tx.vat_amount), (7, tx.total_amount)):
            cell = ws.cell(row=row_idx, column=col, value=amount)
            cell.number_format = '"₩"#,##0'
        ws.cell(row=row_idx, column=8, value=tx.tax_invoice_no or "-")
        ws.cell(row=row_idx, column=9, value=tx.memo or "-")

        if tx.type == TransactionType.sales:
            total_sales += tx.total_amount
        else:
            total_purchase += tx.total_amount
        row_idx += 1

    summary_rows = [
        ("합계 매출", total_sales),
        ("합계 매입", total_purchase),
        ("차액 (매출-매입)", total_sales - total_purchase),
    ]
    row_idx += 1
    for label, value in summary_rows:
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=row_idx, column=7, value=value)
        value_cell.font = Font(bold=True)
        value_cell.number_format = '"₩"#,##0'
        row_idx += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
