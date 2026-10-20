"""품의서를 엑셀(.xlsx)로 내보낸다 (승인된 품의서의 내용·결재 현황·첨부 목록을 한 장에 정리). 저장해 두지 않고 요청할 때마다 만든다."""
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.approval_step import ApprovalStepStatus
from app.models.proposal import Proposal

STEP_STATUS_LABELS = {
    ApprovalStepStatus.pending: "대기",
    ApprovalStepStatus.approved: "승인",
    ApprovalStepStatus.rejected: "반려",
}
TITLE_LABELS = {"dept_head": "부서장", "team_lead": "팀장", "ceo": "대표"}


def _title_label(title) -> str:
    value = getattr(title, "value", title)
    return TITLE_LABELS.get(str(value), str(value))


def generate_proposal_excel(p: Proposal) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "품의서"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    label_fill = PatternFill(start_color="F2F2F3", end_color="F2F2F3", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.merge_cells("A1:D1")
    ws["A1"] = f"{p.kind}품의서" if p.kind else "품의서"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")

    rows = [
        ("문서번호", p.doc_no),
        ("기안일자", p.issue_date.isoformat()),
        ("부서", p.department_name),
        ("기안자", p.creator.name if p.creator else ""),
        ("제목", p.title),
        ("진행 내용", f"아래와 같이 {p.topic}을(를) 진행하고자 합니다."),
        ("내용", p.content),
    ]
    row = 3
    for label, value in rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=1).fill = label_fill
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=2, value=value)
        cell.alignment = wrap
        if label == "내용":
            ws.row_dimensions[row].height = max(60, 15 * (str(value).count("\n") + 1 + len(str(value)) // 45))
        row += 1

    row += 1
    for col, text in enumerate(["단계", "직책", "결재자", "상태 / 일시"], start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    row += 1
    if p.is_final_decision or not p.steps:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row=row, column=1, value="전결" if p.is_final_decision else "결재선 없음").alignment = Alignment(horizontal="center")
        row += 1
    for step in p.steps:
        decided = step.decided_at.strftime("%Y-%m-%d %H:%M") if step.decided_at else ""
        ws.cell(row=row, column=1, value=step.step_order).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2, value=_title_label(step.title)).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=3, value=step.approver.name if step.approver else "").alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=4, value=f"{STEP_STATUS_LABELS.get(step.status, step.status)} {decided}".strip())
        row += 1

    if p.attachments:
        row += 1
        ws.cell(row=row, column=1, value="첨부").font = Font(bold=True)
        ws.cell(row=row, column=1).fill = label_fill
        for a in p.attachments:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            ws.cell(row=row, column=2, value=a.label if a.kind.value == "file" else f"{a.label} ({a.url})")
            row += 1

    for col, width in zip(range(1, 5), (12, 16, 18, 40)):
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
