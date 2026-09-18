from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import settings
from app.models.payment import Payment
from app.services.payment_voucher_pdf import TYPE_LABELS, method_detail_text, proof_text


def generate_payment_voucher_excel(payment: Payment) -> bytes:
    """입출금 내역 1건을 전표 형태의 엑셀(.xlsx) 파일로 생성한다. payment_voucher_pdf.py와 항목 구성을 동일하게 맞춘다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "전표"

    type_label = TYPE_LABELS[payment.type]

    rows: list[tuple[str, str | int]] = [
        ("구분", type_label),
        ("날짜", payment.payment_date.isoformat()),
        ("분류", payment.category),
        ("항목", payment.description),
        ("거래처", payment.client.name if payment.client else "-"),
        ("결제수단", method_detail_text(payment)),
        ("증빙발행", proof_text(payment)),
        ("금액", payment.amount),
        ("메모", payment.memo or "-"),
        ("작성자", payment.creator.name if payment.creator else "-"),
    ]

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 50

    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = f"{settings.company_name} {type_label} 전표"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    label_fill = PatternFill(start_color="F2F2F3", end_color="F2F2F3", fill_type="solid")
    for i, (label, value) in enumerate(rows, start=2):
        label_cell = ws.cell(row=i, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.fill = label_fill
        label_cell.alignment = Alignment(vertical="center")

        value_cell = ws.cell(row=i, column=2, value=value)
        value_cell.alignment = Alignment(vertical="center", wrap_text=True)
        if label == "금액":
            value_cell.number_format = '"₩"#,##0'
        ws.row_dimensions[i].height = 20

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
