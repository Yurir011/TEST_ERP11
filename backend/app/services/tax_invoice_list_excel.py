from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.project_document import ProjectDocument, ProjectDocumentStatus

HEADERS = ["문서번호", "발행일", "거래처", "담당자", "공급가액", "부가세", "합계", "상태", "팝빌발행", "국세청승인번호"]
COLUMN_WIDTHS = [16, 12, 20, 14, 14, 12, 14, 10, 10, 22]

STATUS_LABELS = {
    ProjectDocumentStatus.draft: "초안",
    ProjectDocumentStatus.pending: "결재 대기",
    ProjectDocumentStatus.approved: "승인됨",
    ProjectDocumentStatus.rejected: "반려됨",
}


def _items_total(doc: ProjectDocument) -> int:
    return sum(item.quantity * item.unit_price for item in doc.items)


def _write_sheet(ws, docs: list[ProjectDocument]) -> int:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    total = 0
    row_idx = 2
    for doc in docs:
        subtotal = _items_total(doc)
        vat = round(subtotal * 0.1)
        grand_total = subtotal + vat
        total += grand_total

        ws.cell(row=row_idx, column=1, value=doc.doc_no or "-")
        ws.cell(row=row_idx, column=2, value=doc.issue_date.isoformat())
        ws.cell(row=row_idx, column=3, value=doc.client_name)
        ws.cell(row=row_idx, column=4, value=doc.manager_name or "-")
        for col, amount in ((5, subtotal), (6, vat), (7, grand_total)):
            cell = ws.cell(row=row_idx, column=col, value=amount)
            cell.number_format = '"₩"#,##0'
        ws.cell(row=row_idx, column=8, value=STATUS_LABELS[doc.status])
        ws.cell(row=row_idx, column=9, value="발행완료" if doc.popbill_issued_at else "미발행")
        ws.cell(row=row_idx, column=10, value=doc.popbill_nts_confirm_num or "-")
        row_idx += 1

    if docs:
        label_cell = ws.cell(row=row_idx + 1, column=6, value="합계")
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=row_idx + 1, column=7, value=total)
        value_cell.font = Font(bold=True)
        value_cell.number_format = '"₩"#,##0'

    return total


def generate_tax_invoice_list_excel(
    billing_docs: list[ProjectDocument],
    receipt_docs: list[ProjectDocument],
    purchase_total: int,
    purchase_count: int,
    period_label: str,
) -> bytes:
    """(검색/기간 필터가 적용된) 세금계산서 목록을 청구/영수/개요 3개 탭으로 내보낸다.
    개요 탭의 매입 합계는 매입매출관리(Transactions)의 매입 거래 중 세금계산서번호가 입력된 건을 집계한 값이다."""
    wb = openpyxl.Workbook()

    ws_billing = wb.active
    ws_billing.title = "청구"
    billing_total = _write_sheet(ws_billing, billing_docs)

    ws_receipt = wb.create_sheet("영수")
    receipt_total = _write_sheet(ws_receipt, receipt_docs)

    ws_summary = wb.create_sheet("개요")
    ws_summary.column_dimensions["A"].width = 16
    ws_summary.column_dimensions["B"].width = 14
    ws_summary.column_dimensions["C"].width = 18

    title_cell = ws_summary.cell(row=1, column=1, value=f"세금계산서 개요 ({period_label})")
    title_cell.font = Font(bold=True, size=13)
    ws_summary.merge_cells("A1:C1")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, header in enumerate(["구분", "건수", "금액"], start=1):
        cell = ws_summary.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sales_count = len(billing_docs) + len(receipt_docs)
    sales_total = billing_total + receipt_total
    unissued_docs = [d for d in billing_docs + receipt_docs if not d.popbill_issued_at]
    unissued_total = sum(_items_total(d) + round(_items_total(d) * 0.1) for d in unissued_docs)

    summary_rows = [
        ("매출 (세금계산서 발행 합계)", sales_count, sales_total),
        ("매입 (매입매출관리 연동)", purchase_count, purchase_total),
        ("미결 (팝빌 미발행)", len(unissued_docs), unissued_total),
    ]
    row_idx = 4
    for label, count, amount in summary_rows:
        ws_summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=row_idx, column=2, value=count)
        amount_cell = ws_summary.cell(row=row_idx, column=3, value=amount)
        amount_cell.number_format = '"₩"#,##0'
        row_idx += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
