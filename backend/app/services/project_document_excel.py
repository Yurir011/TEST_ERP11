import os
from io import BytesIO

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import settings
from app.logging_config import get_logger
from app.models.client import Client
from app.models.project import Project
from app.models.project_document import ProjectDocType, ProjectDocument, ProjectDocumentStatus, calc_amount
from app.services.certificate_pdf import STAMP_PATH

logger = get_logger("ProjectDocumentExcel")

DOC_TITLES = {
    ProjectDocType.quotation: "견적서",
    ProjectDocType.statement: "거래명세서",
    ProjectDocType.tax_invoice: "세금계산서",
}

ITEM_HEADERS = ["내용", "수량", "단가", "금액", "비고"]


def generate_project_document_excel(doc: ProjectDocument, project: Project | None, client: Client | None) -> bytes:
    """승인 완료된 견적서/거래명세서/세금계산서를 회계 처리에 바로 쓸 수 있는 엑셀(.xlsx)로 내보낸다.
    project_document_pdf.py와 항목 구성을 동일하게 맞춘다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = DOC_TITLES[doc.doc_type]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    label_fill = PatternFill(start_color="F2F2F3", end_color="F2F2F3", fill_type="solid")

    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"{DOC_TITLES[doc.doc_type]} ({doc.issue_date.isoformat()})"
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    meta_rows = [
        ("프로젝트", doc.project_display_name),
        ("공급자", settings.company_name),
        ("공급자 대표자", settings.company_ceo_name or "-"),
        ("거래처", doc.client_name),
        ("거래처 대표자", (client.ceo_name if client else None) or "-"),
        ("담당자", doc.manager_name or "-"),
    ]
    row_idx = 3
    for label, value in meta_rows:
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.fill = label_fill
        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=5)
        ws.cell(row=row_idx, column=2, value=value)
        row_idx += 1

    row_idx += 1
    item_header_row = row_idx
    for col, header in enumerate(ITEM_HEADERS, start=1):
        cell = ws.cell(row=item_header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row_idx += 1

    subtotal = 0
    for item in doc.items:
        amount = calc_amount(item.quantity, item.unit_price)
        subtotal += amount
        ws.cell(row=row_idx, column=1, value=item.content)
        ws.cell(row=row_idx, column=2, value=item.quantity)
        unit_price_cell = ws.cell(row=row_idx, column=3, value=item.unit_price)
        unit_price_cell.number_format = '"₩"#,##0' if float(item.unit_price).is_integer() else '"₩"#,##0.00'
        amount_cell = ws.cell(row=row_idx, column=4, value=amount)
        amount_cell.number_format = '"₩"#,##0'
        ws.cell(row=row_idx, column=5, value=item.note or "-")
        row_idx += 1

    vat = round(subtotal * 0.1)
    total = subtotal + vat
    totals_rows = [("공급가액", subtotal), ("부가세(10%)", vat), ("합계", total)]
    row_idx += 1
    for label, value in totals_rows:
        label_cell = ws.cell(row=row_idx, column=4, value=label)
        label_cell.font = Font(bold=True)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=row_idx, column=5, value=value)
        value_cell.font = Font(bold=True)
        value_cell.number_format = '"₩"#,##0'
        row_idx += 1

    row_idx += 1
    approver_name = doc.approver.name if doc.approver else "-"
    if doc.is_final_decision:
        approver_name += " (전결)"
    reviewed_date = doc.reviewed_at.date().isoformat() if doc.reviewed_at else "-"
    ws.cell(row=row_idx, column=1, value=f"결재자 : {approver_name}").font = Font(bold=True)
    ws.cell(row=row_idx, column=3, value=f"승인일 : {reviewed_date}")

    if doc.status == ProjectDocumentStatus.approved and os.path.exists(STAMP_PATH):
        try:
            stamp_img = XLImage(STAMP_PATH)
            stamp_img.width = 60
            stamp_img.height = 60
            ws.add_image(stamp_img, f"D{row_idx - 1}")
        except Exception as err:
            logger.debug(f"[ProjectDocumentExcel] 직인 이미지 삽입 실패, 이미지 없이 계속 진행: {err}")

    for col, width in enumerate([26, 10, 14, 14, 22], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
