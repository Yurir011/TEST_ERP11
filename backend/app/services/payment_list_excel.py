from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.payment import Payment, PaymentType
from app.services.payment_voucher_pdf import TYPE_LABELS, method_detail_text, proof_text

HEADERS = ["구분", "날짜", "분류", "항목", "거래처", "결제수단", "증빙발행", "금액", "메모", "작성자"]
COLUMN_WIDTHS = [8, 12, 12, 26, 16, 18, 14, 14, 24, 10]


def generate_payment_list_excel(payments: list[Payment]) -> bytes:
    """(검색/기간 필터가 적용된) 입출금 내역 목록을 엑셀 한 장으로 내보낸다. 맨 아래에 입금/출금/순증감 합계를 덧붙인다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "입출금내역"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for col, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    total_deposit = 0
    total_withdrawal = 0
    row_idx = 2
    for p in payments:
        ws.cell(row=row_idx, column=1, value=TYPE_LABELS[p.type])
        ws.cell(row=row_idx, column=2, value=p.payment_date.isoformat())
        ws.cell(row=row_idx, column=3, value=p.category)
        ws.cell(row=row_idx, column=4, value=p.description)
        ws.cell(row=row_idx, column=5, value=p.client.name if p.client else "-")
        ws.cell(row=row_idx, column=6, value=method_detail_text(p))
        ws.cell(row=row_idx, column=7, value=proof_text(p))
        amount_cell = ws.cell(row=row_idx, column=8, value=p.amount)
        amount_cell.number_format = '"₩"#,##0'
        ws.cell(row=row_idx, column=9, value=p.memo or "-")
        ws.cell(row=row_idx, column=10, value=p.creator.name if p.creator else "-")

        if p.type == PaymentType.deposit:
            total_deposit += p.amount
        else:
            total_withdrawal += p.amount
        row_idx += 1

    summary_rows = [
        ("합계 입금", total_deposit),
        ("합계 출금", total_withdrawal),
        ("순증감", total_deposit - total_withdrawal),
    ]
    row_idx += 1
    for label, value in summary_rows:
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        label_cell.alignment = Alignment(horizontal="right")
        value_cell = ws.cell(row=row_idx, column=8, value=value)
        value_cell.font = Font(bold=True)
        value_cell.number_format = '"₩"#,##0'
        row_idx += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
