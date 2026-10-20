"""견적서·거래명세서 엑셀: 긴 내용이 칸에 맞춰 줄바꿈되고 행 높이가 함께 늘어난다."""
from datetime import date
from types import SimpleNamespace as NS

import openpyxl
import pytest

from app.config import settings
from app.models.project_document import ProjectDocType
from app.services import project_document_template as t


@pytest.fixture(autouse=True)
def _storage_in_tmp(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "project_documents_dir", str(tmp_path / "docs"))


def _doc(kind: ProjectDocType, client: str, manager: str, items):
    return NS(
        id=1,
        doc_type=kind,
        issue_date=date(2026, 10, 20),
        client_name=client,
        manager_name=manager,
        doc_no="1020-01" if kind == ProjectDocType.quotation else None,
        currency="KRW",
        items=items,
    )


LONG_CLIENT = "주식회사 아주아주긴이름의통신장비유통거래처 서울지점"
LONG_MANAGER = "김철수 책임연구원(연구개발 1팀)"
LONG_ITEM = "5G 통신 중계기 RF 모듈 (900MHz/2.1GHz 듀얼밴드, 출력 20W, 방열판 포함 알루미늄 하우징, IP67 방수 규격 적용) 시제품 제작 및 시험 일체"


@pytest.mark.parametrize("kind", [ProjectDocType.quotation, ProjectDocType.statement])
def test_long_client_and_manager_wrap_and_grow_the_row(kind):
    items = [NS(content="짧은 품목", quantity=1, unit_price=1000, note=None)]
    path = t.fill_project_document_template(_doc(kind, LONG_CLIENT, LONG_MANAGER, items))
    ws = openpyxl.load_workbook(path).active
    layout = t.LAYOUTS[kind]
    for ref in (layout.client_cell, layout.manager_cell):
        cell = ws[ref]
        assert cell.alignment.wrap_text is True
        assert not cell.alignment.shrink_to_fit  # 글자를 줄여서 맞추지 않는다
        assert ws.row_dimensions[cell.row].height > 20.25  # 줄 수만큼 행이 높아진다


@pytest.mark.parametrize("kind", [ProjectDocType.quotation, ProjectDocType.statement])
def test_short_client_keeps_the_default_row_height(kind):
    items = [NS(content="짧은 품목", quantity=1, unit_price=1000, note=None)]
    path = t.fill_project_document_template(_doc(kind, "벤치", "홍길동", items))
    ws = openpyxl.load_workbook(path).active
    layout = t.LAYOUTS[kind]
    original = openpyxl.load_workbook(t.TEMPLATE_DIR / layout.template_filename).active
    for ref in (layout.client_cell, layout.manager_cell):
        row = ws[ref].row
        assert (ws.row_dimensions[row].height or 20.25) == (original.row_dimensions[row].height or 20.25)  # 짧으면 높이를 키우지 않는다


@pytest.mark.parametrize("kind", [ProjectDocType.quotation, ProjectDocType.statement])
def test_long_item_content_and_note_wrap_with_taller_row(kind):
    items = [NS(content=LONG_ITEM, quantity=3, unit_price=1500000, note="납기: 발주 후 4주 이내, 부품 수급 상황에 따라 변동 가능, 검수 후 잔금 지급")]
    path = t.fill_project_document_template(_doc(kind, "벤치", "홍길동", items))
    ws = openpyxl.load_workbook(path).active
    layout = t.LAYOUTS[kind]
    row = layout.item_start_row
    assert ws[f"B{row}"].alignment.wrap_text is True and ws[f"I{row}"].alignment.wrap_text is True
    assert ws.row_dimensions[row].height > 60  # 긴 내용은 여러 줄이 되도록 행이 충분히 높아진다
    assert (ws.row_dimensions[row + 1].height or 20.25) <= 20.25 + 0.01  # 내용이 없는 행은 그대로
