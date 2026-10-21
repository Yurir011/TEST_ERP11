"""품의서 결재 특성화 테스트 — 결재 단계 테이블 통합(09 문서) 전후에 모두 통과해야 한다.

결재 단계가 어느 테이블에 저장되는지와 무관하게 API 응답과 동작만 확인한다.
PDF 파일 생성은 대상이 아니므로 만들지 않도록 대체한다.
"""
import pytest
from sqlalchemy import text

from app.models.proposal import Proposal
from app.models.user import JobTitle, UserRole
from app.services import proposal_files


@pytest.fixture(autouse=True)
def _no_pdf(monkeypatch):
    monkeypatch.setattr(proposal_files, "regenerate_pdf", lambda p: None)


# ── API 경로 목록 (proposals.py 분할 전후 같아야 한다) ─────────────────

EXPECTED_OPERATIONS = [
    ('DELETE', '/api/proposals/{proposal_id}', 'delete_proposal_api_proposals__proposal_id__delete', ['204', '422']),
    ('DELETE', '/api/proposals/{proposal_id}/attachments/{attachment_id}', 'delete_attachment_api_proposals__proposal_id__attachments__attachment_id__delete', ['200', '422']),
    ('GET', '/api/proposals', 'list_proposals_api_proposals_get', ['200', '422']),
    ('GET', '/api/proposals/{proposal_id}/attachments/{attachment_id}/file', 'download_attachment_file_api_proposals__proposal_id__attachments__attachment_id__file_get', ['200', '422']),
    ('GET', '/api/proposals/{proposal_id}/excel', 'download_proposal_excel_api_proposals__proposal_id__excel_get', ['200', '422']),
    ('GET', '/api/proposals/{proposal_id}/pdf', 'download_proposal_pdf_api_proposals__proposal_id__pdf_get', ['200', '422']),
    ('POST', '/api/proposals', 'create_proposal_api_proposals_post', ['201', '422']),
    ('POST', '/api/proposals/{proposal_id}/attachments/file', 'upload_attachment_file_api_proposals__proposal_id__attachments_file_post', ['201', '422']),
    ('POST', '/api/proposals/{proposal_id}/attachments/link', 'add_attachment_link_api_proposals__proposal_id__attachments_link_post', ['201', '422']),
    ('POST', '/api/proposals/{proposal_id}/request-delete', 'request_delete_proposal_api_proposals__proposal_id__request_delete_post', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}', 'update_proposal_api_proposals__proposal_id__put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/order-payment', 'set_order_payment_api_proposals__proposal_id__order_payment_put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/order-sent', 'set_order_sent_api_proposals__proposal_id__order_sent_put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/received', 'set_received_api_proposals__proposal_id__received_put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/attachments/{attachment_id}/kind', 'set_attachment_kind_api_proposals__proposal_id__attachments__attachment_id__kind_put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/approve', 'approve_proposal_api_proposals__proposal_id__approve_put', ['200', '422']),
    ('PUT', '/api/proposals/{proposal_id}/reject', 'reject_proposal_api_proposals__proposal_id__reject_put', ['200', '422']),
]


def test_api_operations_are_unchanged():
    from app.main import app

    actual = [
        (method.upper(), path, op["operationId"], sorted(op["responses"]))
        for path, ops in app.openapi()["paths"].items()
        if path.startswith("/api/proposals")
        for method, op in ops.items()
    ]
    assert sorted(actual) == sorted(EXPECTED_OPERATIONS)


@pytest.fixture()
def staff(make_user):
    return {
        "writer": make_user(name="기안자"),
        "other": make_user(name="다른직원"),
        "dept_head": make_user(title=JobTitle.dept_head, name="부서장"),
        "team_lead": make_user(title=JobTitle.team_lead, name="팀장"),
        "ceo": make_user(role=UserRole.admin, title=JobTitle.ceo, name="대표"),
    }


def _create(client, headers, approver_ids=(), end_title="team_lead", final=False):
    res = client.post(
        "/api/proposals",
        headers=headers,
        json={
            "title": "장비 구매", "topic": "계측기", "content": "구매 요청", "issue_date": "2026-10-16",
            "is_final_decision": final, "end_title": None if final else end_title, "approver_ids": list(approver_ids),
        },
    )
    assert res.status_code == 201, res.text
    return res.json()


def _approve(client, headers, pid):
    return client.put(f"/api/proposals/{pid}/approve", headers=headers)


def _step_rows(db, proposal_id) -> int:
    """결재 단계 행 수 (공용 테이블. 옛 전용 테이블은 리비전 0006 에서 지웠다)."""
    return db.execute(
        text("SELECT count(*) FROM approval_steps WHERE target_type::text = 'proposal' AND target_id = :i"), {"i": proposal_id}
    ).scalar()


def test_two_step_approval_and_steps_shape(client, staff, auth_headers):
    p = _create(client, auth_headers(staff["writer"]), [staff["dept_head"].id, staff["team_lead"].id])
    assert p["status"] == "pending" and p["current_step"] == 1
    assert [(s["step_order"], s["title"], s["approver_id"], s["approver_name"], s["status"]) for s in p["steps"]] == [
        (1, "dept_head", staff["dept_head"].id, "부서장", "pending"),
        (2, "team_lead", staff["team_lead"].id, "팀장", "pending"),
    ]

    assert _approve(client, auth_headers(staff["team_lead"]), p["id"]).status_code == 403  # 아직 차례 아님

    first = _approve(client, auth_headers(staff["dept_head"]), p["id"])
    assert first.status_code == 200, first.text
    body = first.json()
    assert body["status"] == "pending" and body["current_step"] == 2
    assert body["steps"][0]["status"] == "approved" and body["steps"][0]["decided_at"] is not None

    last = _approve(client, auth_headers(staff["team_lead"]), p["id"]).json()
    assert last["status"] == "approved"
    assert [s["status"] for s in last["steps"]] == ["approved", "approved"]

    again = _approve(client, auth_headers(staff["team_lead"]), p["id"])
    assert again.status_code == 400


def test_reject(client, staff, auth_headers):
    p = _create(client, auth_headers(staff["writer"]), [staff["dept_head"].id, staff["team_lead"].id])
    res = client.put(f"/api/proposals/{p['id']}/reject", headers=auth_headers(staff["dept_head"]), json={"reason": "예산 초과"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["status"] == "rejected" and body["reject_reason"] == "예산 초과"
    assert [s["status"] for s in body["steps"]] == ["rejected", "pending"]
    assert _approve(client, auth_headers(staff["team_lead"]), p["id"]).status_code == 400


def test_final_decision_is_approved_immediately_without_steps(client, staff, auth_headers):
    p = _create(client, auth_headers(staff["writer"]), final=True)
    assert p["status"] == "approved" and p["is_final_decision"] is True and p["steps"] == []


def test_self_approval_skips_own_steps(client, staff, auth_headers):
    # 부서장이 기안하고 자기 자신이 1단계 결재자 → 1단계는 즉시 승인, 팀장 차례
    p = _create(client, auth_headers(staff["dept_head"]), [staff["dept_head"].id, staff["team_lead"].id])
    assert p["status"] == "pending" and p["current_step"] == 2
    assert [s["status"] for s in p["steps"]] == ["approved", "pending"]

    only_self = _create(client, auth_headers(staff["dept_head"]), [staff["dept_head"].id], end_title="dept_head")
    assert only_self["status"] == "approved"


def test_my_inbox_lists_only_current_step_approver(client, staff, auth_headers):
    p = _create(client, auth_headers(staff["writer"]), [staff["dept_head"].id, staff["team_lead"].id])
    mine = lambda u: [x["id"] for x in client.get("/api/proposals?status=pending&approver_mine=true", headers=auth_headers(u)).json()]  # noqa: E731
    assert p["id"] in mine(staff["dept_head"])
    assert p["id"] not in mine(staff["team_lead"])
    _approve(client, auth_headers(staff["dept_head"]), p["id"])
    assert p["id"] not in mine(staff["dept_head"])
    assert p["id"] in mine(staff["team_lead"])


def test_notification_shows_my_approval_done(client, staff, auth_headers):
    p = _create(client, auth_headers(staff["writer"]), [staff["dept_head"].id, staff["team_lead"].id])
    h = auth_headers(staff["dept_head"])

    def my_request_notice():
        items = client.get("/api/notifications", headers=h).json()
        return next(n for n in items if n["title"].endswith("결재 요청") and f"open={p['id']}" in (n["link"] or ""))

    assert my_request_notice()["my_approval_done"] is False
    _approve(client, h, p["id"])
    assert my_request_notice()["my_approval_done"] is True


def test_delete_request_approved_deletes_proposal_and_steps(client, db, staff, auth_headers):
    writer_h = auth_headers(staff["writer"])
    p = _create(client, writer_h, [staff["dept_head"].id], end_title="dept_head")
    _approve(client, auth_headers(staff["dept_head"]), p["id"])
    res = client.post(
        f"/api/proposals/{p['id']}/request-delete", headers=writer_h, json={"end_title": "dept_head", "approver_ids": [staff["dept_head"].id]}
    )
    assert res.status_code == 200, res.text
    assert res.json()["delete_requested"] is True and res.json()["status"] == "pending"

    assert _approve(client, auth_headers(staff["dept_head"]), p["id"]).status_code == 200
    assert db.query(Proposal).filter(Proposal.id == p["id"]).count() == 0
    assert _step_rows(db, p["id"]) == 0  # 결재 단계도 남지 않는다


def test_delete_request_rejected_restores_previous_approval(client, staff, auth_headers):
    writer_h = auth_headers(staff["writer"])
    p = _create(client, writer_h, [staff["dept_head"].id, staff["team_lead"].id])
    _approve(client, auth_headers(staff["dept_head"]), p["id"])
    approved = _approve(client, auth_headers(staff["team_lead"]), p["id"]).json()

    client.post(f"/api/proposals/{p['id']}/request-delete", headers=writer_h, json={"end_title": "dept_head", "approver_ids": [staff["dept_head"].id]})
    res = client.put(f"/api/proposals/{p['id']}/reject", headers=auth_headers(staff["dept_head"]), json={"reason": "보관 필요"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["status"] == "approved" and body["delete_requested"] is False
    # 삭제 요청 전의 결재선(2단계, 둘 다 승인)으로 돌아간다
    assert [(s["step_order"], s["approver_id"], s["status"]) for s in body["steps"]] == [
        (s["step_order"], s["approver_id"], s["status"]) for s in approved["steps"]
    ]


def test_delete_pending_proposal_removes_steps(client, db, staff, auth_headers):
    writer_h = auth_headers(staff["writer"])
    p = _create(client, writer_h, [staff["dept_head"].id, staff["team_lead"].id])
    assert _step_rows(db, p["id"]) == 2
    assert client.delete(f"/api/proposals/{p['id']}", headers=writer_h).status_code == 204
    assert _step_rows(db, p["id"]) == 0


def test_delete_rules(client, staff, auth_headers):
    writer_h = auth_headers(staff["writer"])
    p = _create(client, writer_h, [staff["dept_head"].id], end_title="dept_head")
    assert client.delete(f"/api/proposals/{p['id']}", headers=auth_headers(staff["other"])).status_code == 403  # 작성자만
    _approve(client, auth_headers(staff["dept_head"]), p["id"])
    assert client.delete(f"/api/proposals/{p['id']}", headers=writer_h).status_code == 400  # 승인된 품의서는 삭제 결재 필요


def test_attachment_file_limit_is_seven(client, staff, auth_headers, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "proposal_attachments_dir", str(tmp_path / "att"))
    h = auth_headers(staff["writer"])
    p = _create(client, h, final=True)
    pdf = b"%PDF-1.7\n" + b"0" * 64
    for i in range(7):
        r = client.post(f"/api/proposals/{p['id']}/attachments/file", files={"file": (f"f{i}.pdf", pdf, "application/pdf")}, headers=h)
        assert r.status_code == 201, r.text
    over = client.post(f"/api/proposals/{p['id']}/attachments/file", files={"file": ("f8.pdf", pdf, "application/pdf")}, headers=h)
    assert over.status_code == 400 and "7개" in over.json()["detail"]
    # 링크는 파일 수에 세지 않는다
    assert client.post(f"/api/proposals/{p['id']}/attachments/link", json={"url": "https://example.com"}, headers=h).status_code == 201
    # 허용하지 않는 형식
    bad = client.post(f"/api/proposals/{p['id']}/attachments/file", files={"file": ("a.zip", b"PK\x03\x04", "application/zip")}, headers=h)
    assert bad.status_code == 400


def test_proposal_excel_download(client, staff, auth_headers):
    """승인된 품의서는 엑셀로 내려받을 수 있고, 승인 전에는 막힌다."""
    import io

    import openpyxl

    h = auth_headers(staff["writer"])
    approved = _create(client, h, final=True)
    res = client.get(f"/api/proposals/{approved['id']}/excel", headers=h)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "filename*=UTF-8''" in res.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert approved["doc_no"] in cells and "장비 구매" in cells and "전결" in cells

    pending = _create(client, h, [staff["dept_head"].id], end_title="dept_head")
    assert client.get(f"/api/proposals/{pending['id']}/excel", headers=h).status_code == 403


def test_pdf_download_regenerates_with_current_layout(client, staff, auth_headers, tmp_path, monkeypatch):
    """내려받을 때마다 현재 양식으로 PDF 를 다시 만든다 (옛 양식으로 저장된 문서도 새 모양으로 나온다)."""
    calls = []

    def fake_regenerate(p):
        calls.append(p.id)
        path = tmp_path / f"{p.id}.pdf"
        path.write_bytes(b"%PDF-1.7\nnew layout")
        p.file_path = str(path)

    monkeypatch.setattr(proposal_files, "regenerate_pdf", fake_regenerate)
    h = auth_headers(staff["writer"])
    p = _create(client, h, final=True)
    calls.clear()
    res = client.get(f"/api/proposals/{p['id']}/pdf", headers=h)
    assert res.status_code == 200 and res.content.endswith(b"new layout")
    assert calls == [p["id"]]


def test_pdf_download_falls_back_to_stored_file_when_regeneration_fails(client, db, staff, auth_headers, tmp_path, monkeypatch):
    h = auth_headers(staff["writer"])
    p = _create(client, h, final=True)
    stored = tmp_path / "old.pdf"
    stored.write_bytes(b"%PDF-1.7\nold file")
    db.get(Proposal, p["id"]).file_path = str(stored)
    db.commit()  # 저장된 파일 경로를 확정해 둔다 (다시 만들기 실패 때의 rollback 이 이 값을 되돌리지 않게)

    def boom(_p):
        raise RuntimeError("font missing")

    monkeypatch.setattr(proposal_files, "regenerate_pdf", boom)
    res = client.get(f"/api/proposals/{p['id']}/pdf", headers=h)
    assert res.status_code == 200 and res.content.endswith(b"old file")


def test_attachment_doc_kind_upload_and_change(client, staff, auth_headers, tmp_path, monkeypatch):
    """첨부 파일에 서류 종류(거래명세서·세금계산서)를 정해 올리고, 나중에 바꿀 수 있다."""
    from app.config import settings

    monkeypatch.setattr(settings, "proposal_attachments_dir", str(tmp_path / "att"))
    h = auth_headers(staff["writer"])
    p = _create(client, h, final=True)
    pdf = b"%PDF-1.7\n" + b"0" * 64

    def upload(name, **data):
        return client.post(f"/api/proposals/{p['id']}/attachments/file", files={"file": (name, pdf, "application/pdf")}, data=data, headers=h)

    statement = upload("거래명세서.pdf", doc_kind="statement")
    assert statement.status_code == 201, statement.text
    plain = upload("기타.pdf")  # 종류를 안 정하면 기타(None)
    other = upload("기타2.pdf", doc_kind="other")
    assert plain.status_code == 201 and other.status_code == 201
    kinds = {a["label"]: a["doc_kind"] for a in other.json()["attachments"]}
    assert kinds == {"거래명세서.pdf": "statement", "기타.pdf": None, "기타2.pdf": None}

    assert upload("x.pdf", doc_kind="receipt").status_code == 400  # 알 수 없는 종류

    att_id = next(a["id"] for a in other.json()["attachments"] if a["label"] == "기타.pdf")
    changed = client.put(f"/api/proposals/{p['id']}/attachments/{att_id}/kind", json={"doc_kind": "tax_invoice"}, headers=h)
    assert changed.status_code == 200
    assert {a["label"]: a["doc_kind"] for a in changed.json()["attachments"]}["기타.pdf"] == "tax_invoice"
    cleared = client.put(f"/api/proposals/{p['id']}/attachments/{att_id}/kind", json={"doc_kind": None}, headers=h)
    assert {a["label"]: a["doc_kind"] for a in cleared.json()["attachments"]}["기타.pdf"] is None
    assert client.put(f"/api/proposals/{p['id']}/attachments/{att_id}/kind", json={"doc_kind": "bogus"}, headers=h).status_code == 400


def test_attachment_doc_kind_permissions_and_links(client, staff, auth_headers, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "proposal_attachments_dir", str(tmp_path / "att"))
    h = auth_headers(staff["writer"])
    p = _create(client, h, final=True)
    pdf = b"%PDF-1.7\n" + b"0" * 64
    up = client.post(f"/api/proposals/{p['id']}/attachments/file", files={"file": ("a.pdf", pdf, "application/pdf")}, headers=h).json()
    att_id = up["attachments"][0]["id"]
    # 작성자가 아니면 바꿀 수 없다
    assert client.put(f"/api/proposals/{p['id']}/attachments/{att_id}/kind", json={"doc_kind": "statement"}, headers=auth_headers(staff["other"])).status_code == 403
    # 링크 첨부는 서류 종류를 정할 수 없다
    link = client.post(f"/api/proposals/{p['id']}/attachments/link", json={"url": "https://example.com"}, headers=h).json()
    link_id = next(a["id"] for a in link["attachments"] if a["kind"] == "link")
    assert client.put(f"/api/proposals/{p['id']}/attachments/{link_id}/kind", json={"doc_kind": "statement"}, headers=h).status_code == 400
    assert client.put(f"/api/proposals/{p['id']}/attachments/999999/kind", json={"doc_kind": "statement"}, headers=h).status_code == 404


def test_proposal_vendor_and_amount_saved_edited_and_in_documents(client, db, staff, auth_headers):
    """구매 품의서의 거래처명·금액·통화를 저장·수정하고, PDF·엑셀에도 나온다."""
    import io

    import openpyxl

    h = auth_headers(staff["writer"])
    body = {
        "kind": "구매", "title": "장비 구매", "topic": "계측기", "content": "구매 요청", "issue_date": "2026-10-16",
        "vendor_name": " 대성RF ", "amount": 1234567, "currency": "KRW", "is_final_decision": True,
    }
    created = client.post("/api/proposals", headers=h, json=body)
    assert created.status_code == 201, created.text
    p = created.json()
    assert (p["vendor_name"], p["amount"], p["currency"]) == ("대성RF", 1234567, "KRW")

    # 비워서 올리면 거래처·금액 없음 (기존 품의서와 같은 모양)
    blank = client.post("/api/proposals", headers=h, json={**body, "vendor_name": "", "amount": None}).json()
    assert blank["vendor_name"] is None and blank["amount"] is None
    assert client.post("/api/proposals", headers=h, json={**body, "amount": -1}).status_code == 422

    xlsx = client.get(f"/api/proposals/{p['id']}/excel", headers=h)
    assert xlsx.status_code == 200
    cells = [c.value for row in openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows() for c in row]
    assert "대성RF" in cells and "₩1,234,567" in cells

    from app.services.proposal_pdf import format_proposal_amount

    assert format_proposal_amount(None, "KRW") == "-"
    assert format_proposal_amount(12.5, "USD") == "$12.50"
    from app.models.proposal import Proposal as ProposalModel
    from app.services.proposal_pdf import generate_proposal_pdf
    import pymupdf

    row = db.get(ProposalModel, p["id"])
    text_out = "".join(page.get_text() for page in pymupdf.open(stream=generate_proposal_pdf(row), filetype="pdf"))
    assert "대성RF" in text_out and "1,234,567" in text_out


def _purchase(client, h, amount=1000000):
    body = {
        "kind": "구매", "title": "소모품 구입", "topic": "소모품 구입", "content": "x", "issue_date": "2026-10-16",
        "vendor_name": "대성RF", "amount": amount, "currency": "KRW", "is_final_decision": True,
    }
    res = client.post("/api/proposals", headers=h, json=body)
    assert res.status_code == 201, res.text
    return res.json()


def _ledger(db, proposal_id):
    from app.models.transaction import Transaction

    db.expire_all()
    return db.query(Transaction).filter(Transaction.proposal_id == proposal_id).all()


def _payments(db, proposal_id):
    from app.models.payment import Payment
    from app.models.proposal import Proposal as ProposalModel

    db.expire_all()
    p = db.get(ProposalModel, proposal_id)
    return [db.get(Payment, p.payment_id)] if p.payment_id else []


def test_order_done_creates_credit_purchase_without_payment(client, db, staff, auth_headers):
    """발주 완료에서 입금·법인카드를 고르지 않으면 외상 매입으로 잡히고 입출금은 생기지 않는다."""
    h = auth_headers(staff["writer"])
    p = _purchase(client, h)
    assert _ledger(db, p["id"]) == []
    res = client.put(f"/api/proposals/{p['id']}/order-sent", headers=h, json={"value": True, "online": True})
    assert res.status_code == 200, res.text
    assert res.json()["payment_method"] is None
    (tx,) = _ledger(db, p["id"])
    assert (tx.type.value, tx.counterparty, tx.supply_amount, tx.vat_amount, tx.total_amount) == ("purchase", "대성RF", 1000000, 100000, 1100000)
    assert tx.settled is False and "외상" in tx.memo
    assert _payments(db, p["id"]) == []
    # 발송 확인을 취소하면 매입도 사라진다
    assert client.put(f"/api/proposals/{p['id']}/order-sent", headers=h, json={"value": False}).status_code == 200
    assert _ledger(db, p["id"]) == []


def test_order_with_card_or_deposit_creates_payment_and_settles(client, db, staff, auth_headers):
    h = auth_headers(staff["writer"])
    p = _purchase(client, h)
    res = client.put(f"/api/proposals/{p['id']}/order-sent", headers=h, json={"value": True, "online": True, "payment_method": "card", "payment_card": "bc"})
    assert res.status_code == 200 and (res.json()["payment_method"], res.json()["payment_card"]) == ("card", "bc")
    (tx,) = _ledger(db, p["id"])
    (pay,) = _payments(db, p["id"])
    assert tx.settled is True and tx.settled_payment_id == pay.id
    assert (pay.type.value, pay.method.value, pay.card_type, pay.amount, pay.card_pending) == ("withdrawal", "corporate_card", "bc", 1100000, False)

    # 입금(계좌)으로 바꾸면 카드 출금을 지우고 계좌이체 출금으로 다시 만든다
    res = client.put(f"/api/proposals/{p['id']}/order-payment", headers=h, json={"payment_method": "deposit", "payment_bank": "ibk"})
    assert res.status_code == 200 and res.json()["payment_card"] is None and res.json()["payment_bank"] == "ibk"
    (pay,) = _payments(db, p["id"])
    assert (pay.method.value, pay.bank_type, pay.card_type) == ("bank_transfer", "ibk", None)
    assert _ledger(db, p["id"])[0].settled is True

    # 외상으로 돌리면 출금이 지워지고 매입이 외상으로 돌아온다
    client.put(f"/api/proposals/{p['id']}/order-payment", headers=h, json={})
    assert _payments(db, p["id"]) == []
    (tx,) = _ledger(db, p["id"])
    assert tx.settled is False and tx.settled_payment_id is None
    # 입금은 계좌, 법인카드는 카드를 골라야 한다
    assert client.put(f"/api/proposals/{p['id']}/order-payment", headers=h, json={"payment_method": "card"}).status_code == 400
    assert client.put(f"/api/proposals/{p['id']}/order-payment", headers=h, json={"payment_method": "deposit"}).status_code == 400


def test_receive_asks_payment_and_applies_to_ledger(client, db, staff, auth_headers):
    """입고 완료 때 대금 결제 여부(외상 / 결제 완료)를 확인해 매입·입출금에 반영한다."""
    h = auth_headers(staff["writer"])
    p = _purchase(client, h)
    client.put(f"/api/proposals/{p['id']}/order-sent", headers=h, json={"value": True, "online": True})
    assert _payments(db, p["id"]) == []

    # 입고하면서 입금 완료로 확인
    res = client.put(f"/api/proposals/{p['id']}/received", headers=h, json={"value": True, "payment_asked": True, "payment_method": "deposit", "payment_bank": "woori"})
    assert res.status_code == 200, res.text
    assert res.json()["payment_method"] == "deposit" and res.json()["received_at"] is not None
    (pay,) = _payments(db, p["id"])
    assert pay.bank_type == "woori" and _ledger(db, p["id"])[0].settled is True

    # 입고 확인을 취소했다 다시 하면서 외상으로 확인하면 출금이 지워진다
    client.put(f"/api/proposals/{p['id']}/received", headers=h, json={"value": False})
    client.put(f"/api/proposals/{p['id']}/received", headers=h, json={"value": True, "payment_asked": True})
    assert _payments(db, p["id"]) == [] and _ledger(db, p["id"])[0].settled is False

    # 결제 여부를 묻지 않고 입고만 하면 기존 결제 상태를 그대로 둔다
    q = _purchase(client, h)
    client.put(f"/api/proposals/{q['id']}/order-sent", headers=h, json={"value": True, "online": True, "payment_method": "card", "payment_card": "kb_kookmin"})
    client.put(f"/api/proposals/{q['id']}/received", headers=h, json={"value": True})
    assert len(_payments(db, q["id"])) == 1


def test_order_payment_needs_order_done_and_writer(client, staff, auth_headers):
    h = auth_headers(staff["writer"])
    p = _purchase(client, h)
    assert client.put(f"/api/proposals/{p['id']}/order-payment", headers=h, json={"payment_method": "deposit"}).status_code == 400
    client.put(f"/api/proposals/{p['id']}/order-sent", headers=h, json={"value": True, "online": True})
    assert client.put(f"/api/proposals/{p['id']}/order-payment", headers=auth_headers(staff["other"]), json={"payment_method": "deposit"}).status_code == 403


def test_proposal_project_selection(client, db, staff, auth_headers):
    """구매 품의서에 프로젝트를 고를 수 있고(기본은 일반 = 없음), 없는 프로젝트는 거절한다."""
    from app.models.client import Client
    from app.models.project import Project

    h = auth_headers(staff["writer"])
    customer = Client(name="고객사", created_by=staff["writer"].id)
    db.add(customer)
    db.flush()
    project = Project(name="신규 모뎀 개발", client_id=customer.id, created_by=staff["writer"].id)
    db.add(project)
    db.commit()
    body = {"kind": "구매", "title": "부품 구입", "topic": "부품", "content": "x", "issue_date": "2026-10-16", "is_final_decision": True}

    general = client.post("/api/proposals", headers=h, json=body).json()
    assert general["project_id"] is None and general["project_name"] is None

    chosen = client.post("/api/proposals", headers=h, json={**body, "project_id": project.id})
    assert chosen.status_code == 201, chosen.text
    assert (chosen.json()["project_id"], chosen.json()["project_name"]) == (project.id, "신규 모뎀 개발")

    assert client.post("/api/proposals", headers=h, json={**body, "project_id": 999999}).status_code == 400
    # 수정으로 일반으로 되돌릴 수 있다
    edited = client.put(f"/api/proposals/{chosen.json()['id']}", headers=h, json={**body, "project_id": None, "is_final_decision": True})
    assert edited.status_code == 200 and edited.json()["project_id"] is None


def test_object_particle_and_pdf_intro_sentence():
    """본문 도입 문장의 조사: 구매를 / 출장을 (받침에 따라)."""
    from app.core.korean import object_particle

    assert object_particle("구매") == "를" and object_particle("출장") == "을" and object_particle("PCB") == "을(를)" and object_particle("") == "을"


def test_doc_no_does_not_collide_after_a_same_day_proposal_is_deleted(client, db, staff, auth_headers):
    """같은 날 품의서를 지운 뒤 새로 써도 문서번호가 겹치지 않는다 (건수 + 1 로 채번하면 이미 있는 번호와 같아져 500 오류가 났다)."""
    h = auth_headers(staff["writer"])
    first, second, third = (_create(client, h, [staff["dept_head"].id, staff["team_lead"].id]) for _ in range(3))
    assert [p["doc_no"][-2:] for p in (first, second, third)] == ["01", "02", "03"]
    assert client.delete(f"/api/proposals/{first['id']}", headers=h).status_code in (200, 204)  # 01 삭제 → 건수는 2, 가장 큰 번호는 03
    fourth = _create(client, h, [staff["dept_head"].id, staff["team_lead"].id])
    assert fourth["doc_no"].endswith("-04") and fourth["doc_no"] not in (second["doc_no"], third["doc_no"])
